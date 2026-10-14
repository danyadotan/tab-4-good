from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PDF_PATH = ROOT / "output/pdf/TAB-Work-Investor-Brief.pdf"
DOCX_PATH = ROOT / "output/docx/TAB-Work-Internal-Product-Brief.docx"
PUBLIC_PDF = ROOT / "public/TAB-Work-Investor-Brief.pdf"

PURPLE = colors.HexColor("#7C3AED")
INK = colors.HexColor("#0A0A0B")
MUTED = colors.HexColor("#667085")
PALE = colors.HexColor("#F4F0FF")
GHOST = colors.HexColor("#E2E8F0")
WHITE = colors.white

FONT_REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
pdfmetrics.registerFont(TTFont("DejaVu", FONT_REG))
pdfmetrics.registerFont(TTFont("DejaVu-Bold", FONT_BOLD))


def wrap(c, text, x, y, width, font="DejaVu", size=8.3, leading=11, color=INK, max_lines=None):
    c.setFont(font, size)
    c.setFillColor(color)
    lines = []
    for paragraph in text.split("\n"):
        if not paragraph:
            lines.append("")
            continue
        words, line = paragraph.split(), ""
        for word in words:
            trial = (line + " " + word).strip()
            if c.stringWidth(trial, font, size) <= width:
                line = trial
            else:
                lines.append(line)
                line = word
        if line:
            lines.append(line)
    if max_lines and len(lines) > max_lines:
        lines = lines[:max_lines]
        while c.stringWidth(lines[-1] + "…", font, size) > width:
            lines[-1] = lines[-1][:-1]
        lines[-1] += "…"
    for row in lines:
        c.drawString(x, y, row)
        y -= leading
    return y


def investor_pdf():
    w, h = A4
    c = canvas.Canvas(str(PDF_PATH), pagesize=A4)
    c.setTitle("TAB@Work — One-page investor brief")
    margin = 36
    c.setFillColor(INK)
    c.rect(0, 0, w, h, fill=1, stroke=0)
    c.setFillColor(PURPLE)
    c.circle(margin + 4, h - 38, 4, fill=1, stroke=0)
    c.setFont("DejaVu-Bold", 10)
    c.setFillColor(WHITE)
    c.drawString(margin + 16, h - 42, "TAB@WORK")
    c.setFont("DejaVu", 7.2)
    c.setFillColor(colors.HexColor("#B8B8C2"))
    c.drawRightString(w - margin, h - 42, "Bridge-iT · Danya Dotan · October 2026")

    c.setFont("DejaVu-Bold", 27)
    c.setFillColor(WHITE)
    c.drawString(margin, h - 91, "The execution layer AI-native")
    c.drawString(margin, h - 123, "organizations are missing.")
    wrap(c, "AI can generate more work than organizations can reliably carry to completion. TAB detects where execution breaks, coordinates the smallest authorized intervention, and recovers continuity across people, systems, and agents.", margin, h - 150, w - 2 * margin, size=9.2, leading=13, color=colors.HexColor("#C9CAD2"), max_lines=3)

    y = h - 205
    col_gap = 12
    col_w = (w - 2 * margin - col_gap) / 2

    def panel(x, y_top, width, height, title, body, accent=False, max_lines=6):
        c.setFillColor(PALE if accent else colors.HexColor("#151519"))
        c.roundRect(x, y_top-height, width, height, 5, fill=1, stroke=0)
        c.setFillColor(PURPLE)
        c.rect(x, y_top-height, 3, height, fill=1, stroke=0)
        c.setFont("DejaVu-Bold", 9)
        c.setFillColor(INK if accent else WHITE)
        c.drawString(x+13, y_top-18, title)
        wrap(c, body, x+13, y_top-35, width-25, size=7.6, leading=10.3, color=colors.HexColor("#414151") if accent else colors.HexColor("#BFC0C8"), max_lines=max_lines)

    panel(margin, y, col_w, 103, "THE THESIS", "The bottleneck is shifting from intelligence to execution. Work fails at seams: unclear ownership, missing permissions, dropped context, duplicate review, and agent-to-human handoffs. TAB owns the recovery loop without replacing source systems.", True)
    panel(margin+col_w+col_gap, y, col_w, 103, "PEOPLE + SYSTEMS", "People retain authority. Systems remain systems of record. AI agents contribute speed. TAB sits across all three: observe workflow signals → classify breakdown → route recovery → verify completion.", False)

    y -= 118
    c.setFont("DejaVu-Bold", 9)
    c.setFillColor(WHITE)
    c.drawString(margin, y, "HOW THE BOUNCER FITS")
    y -= 16
    labels = [("PEOPLE", "judgment · authority"), ("BOUNCER", "hygiene · closure"), ("TAB CORE", "detect · coordinate · recover"), ("SYSTEMS + AI", "work · evidence · actions")]
    box_w = (w - 2*margin - 18) / 4
    for i, (head, sub) in enumerate(labels):
        x = margin + i * (box_w + 6)
        c.setFillColor(PURPLE if i == 1 else colors.HexColor("#202027"))
        c.roundRect(x, y-48, box_w, 48, 4, fill=1, stroke=0)
        c.setFont("DejaVu-Bold", 7.3)
        c.setFillColor(WHITE)
        c.drawCentredString(x+box_w/2, y-18, head)
        c.setFont("DejaVu", 6.5)
        c.setFillColor(colors.HexColor("#D8D8E0"))
        c.drawCentredString(x+box_w/2, y-33, sub)

    y -= 68
    panel(margin, y, col_w, 145, "VALIDATION NOW", "• US design partnership in a trading workflow\n• Call-center research at the earliest operational breakpoints\n• Behavior-first prototypes testing timing, trust, and recovery\n\nCurrent proof is early-stage. The next milestone is repeatable workflow-level evidence in 3–5 live environments.", False, 10)
    panel(margin+col_w+col_gap, y, col_w, 145, "A FOUR-WEEK PILOT", "W1 — map one failing workflow and baseline\nW2 — instrument signals and authority rules\nW3 — run recovery in shadow/live mode\nW4 — compare latency, drops, rework, and closure\n\nNo expansion before the evidence clears the gate.", True, 10)

    y -= 160
    c.setFillColor(WHITE)
    c.setFont("DejaVu-Bold", 9)
    c.drawString(margin, y, "THE OPPORTUNITY")
    y -= 15
    wrap(c, "Design partners get a bounded pilot, direct influence on the recovery model, and an evidence pack for rollout decisions. Investors fund the transition from a working architecture and field validation into repeatable enterprise proof, connector depth, and a defensible execution-friction benchmark.", margin, y, w-2*margin, size=7.7, leading=10.5, color=colors.HexColor("#C9CAD2"), max_lines=4)

    c.setStrokeColor(colors.HexColor("#34343C"))
    c.line(margin, 55, w-margin, 55)
    c.setFont("DejaVu-Bold", 8.3)
    c.setFillColor(WHITE)
    c.drawString(margin, 35, "Design partners · enterprise operators · pre-seed investors")
    c.setFont("DejaVu", 7.5)
    c.setFillColor(colors.HexColor("#B8B8C2"))
    c.drawRightString(w-margin, 35, "danya@dynamicbridge.io")
    c.showPage()
    c.save()
    PUBLIC_PDF.write_bytes(PDF_PATH.read_bytes())


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color="0A0A0B", size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for i, head in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], head, True, "FFFFFF", 8)
        set_cell_shading(table.rows[0].cells[i], "0A0A0B")
        if widths:
            table.rows[0].cells[i].width = Inches(widths[i])
    set_repeat_table_header(table.rows[0])
    for row_i, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], str(value), False, "0A0A0B", 8)
            if row_i % 2:
                set_cell_shading(cells[i], "F7F7FA")
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def internal_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.6)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for name, size, color in [("Title", 30, "0A0A0B"), ("Heading 1", 20, "0A0A0B"), ("Heading 2", 13, "7C3AED"), ("Heading 3", 10.5, "0A0A0B")]:
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)
        styles[name].font.bold = True
        styles[name].paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.space_before = Pt(4)
    styles["Heading 1"].paragraph_format.space_after = Pt(10)
    styles["Heading 2"].paragraph_format.space_before = Pt(8)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "TAB@WORK  /  INTERNAL PRODUCT BRIEF  /  OCTOBER 2026"
        header.style = styles["Caption"]
        header.runs[0].font.color.rgb = RGBColor(124, 58, 237)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Bridge-iT · Confidential  |  ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        footer._p.append(fld)

    p = doc.add_paragraph()
    p.add_run("BRIDGE-iT  /  PRODUCT & GTM").bold = True
    p.runs[0].font.color.rgb = RGBColor(124, 58, 237)
    doc.add_heading("TAB@Work", 0)
    p = doc.add_paragraph("Internal product brief")
    p.style = styles["Subtitle"]
    doc.add_paragraph("The operating model, product architecture, packaging, buyer thesis, and current implementation behind an execution-recovery layer for AI-native organizations.")
    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    set_cell_shading(callout.cell(0, 0), "F4F0FF")
    set_cell_text(callout.cell(0, 0), "STATUS  ·  Working product thesis and commercial test plan. External market figures are sourced; TAB pricing remains a benchmark-derived hypothesis to validate in paid pilots.", True, "4C1D95", 9)
    doc.add_paragraph()
    add_table(doc, ["What TAB is", "What TAB is not"], [
        ("A horizontal detection, coordination, and recovery layer across people, systems, and AI agents.", "A replacement for Jira, Slack, Salesforce, a copilot, or a system of record."),
        ("A workflow-level reliability mechanism with explicit authority and evidence requirements.", "Employee surveillance, sentiment scoring, or an autonomous buyer/approver."),
        ("A way to recover the handoff and verify closure.", "Another dashboard that only reports the stall."),
    ], [3.35, 3.35])
    doc.add_heading("Executive position", 1)
    doc.add_paragraph("AI expands the volume and speed of proposed work. Organizations still rely on fragmented approvals, permissions, handoffs, and human judgment to turn proposals into completed outcomes. TAB addresses that gap: it detects the breakdown, routes the smallest authorized recovery action, preserves context, and closes only when evidence supports completion.")
    doc.add_heading("The wedge", 2)
    add_bullets(doc, [
        "Land on one measurable workflow where execution repeatedly stalls.",
        "Integrate at the workflow seam rather than replacing the tools on either side.",
        "Prove reduction in recovery latency, dropped handoffs, rework, or unresolved exceptions.",
        "Expand only after the workflow has a verified baseline and a credible counterfactual.",
    ])

    doc.add_page_break()
    doc.add_heading("1. Day-to-day: trading operations", 1)
    doc.add_paragraph("Scenario: a trading team needs to act on a time-sensitive opportunity, but quote versions, approval authority, onboarding terms, and delivery evidence are split across email, chat, CRM, and finance.")
    add_table(doc, ["Moment", "Breakdown", "TAB intervention", "Evidence of value"], [
        ("A new quote appears", "The team cannot tell whether it supersedes the prior quote.", "Detect conflicting identifiers and route a source-of-truth check to Procurement.", "Fewer decisions based on stale commercial terms."),
        ("Price exceeds authority", "A manager is asked to approve beyond a documented threshold.", "Coordinate an exception request to the authorized approver; block commitment.", "Zero unauthorized commitments; faster exception routing."),
        ("A delivery promise is repeated", "Start date and calendar basis are missing.", "Recover the contractual dates before presenting a human choice.", "Lower rework and fewer false deadlines."),
        ("Work is complete", "Status is declared without execution evidence.", "Closure Gate requires transaction or downstream completion evidence.", "Verified closure rate, not activity count."),
    ], [1.15, 1.9, 2.35, 1.3])
    doc.add_heading("Best pilot metric", 2)
    doc.add_paragraph("Median recovery latency from detected exception to an authority-complete decision, paired with the percentage of recommendations returned for missing evidence.")
    doc.add_heading("Privacy boundary", 2)
    doc.add_paragraph("Use transaction and workflow state. Do not infer individual performance, emotional state, or intent.")

    doc.add_page_break()
    doc.add_heading("2. Day-to-day: call-center onboarding", 1)
    doc.add_paragraph("Scenario: new representatives enter a high-volume environment with training, identity, system access, schedule, and knowledge steps owned by different teams. Attrition risk often appears first as unresolved setup friction.")
    add_table(doc, ["Signal", "Typical owner", "Recovery action", "Guardrail"], [
        ("Identity created; CRM access missing", "IT / Security", "Create one permission request with the minimum role and due date.", "No credential sharing; no privilege expansion without approval."),
        ("Training complete; live-call readiness unclear", "Enablement / Team lead", "Request a bounded readiness check tied to published criteria.", "No inferred psychological score or hidden ranking."),
        ("Knowledge answer conflicts with policy", "Quality / Compliance", "Hold the script, reconcile the source, and publish the verified answer.", "The policy system remains authoritative."),
        ("Early friction repeats across cohort", "Operations", "Surface aggregate pattern and route root-cause review.", "Cohort metrics only; do not expose individual-level behavior."),
    ], [1.45, 1.25, 2.55, 1.45])
    doc.add_heading("Best pilot metrics", 2)
    add_bullets(doc, [
        "Time from start date to access-complete readiness.",
        "Percentage of onboarding handoffs without an owner after 24 hours.",
        "Repeat contacts caused by unresolved access or policy context.",
        "Early-stage attrition is a lagging measure and should not be claimed as causal during a four-week pilot.",
    ])

    doc.add_page_break()
    doc.add_heading("3. Day-to-day: enterprise AI rollout", 1)
    doc.add_paragraph("Scenario: teams adopt copilots and agents faster than governance can distinguish routine, reversible output from material decisions that require authority.")
    add_table(doc, ["Agent output", "Hygiene Gate", "Closure Gate", "Outcome"], [
        ("Three overlapping recommendations", "Deduplicate and reconcile contradictions.", "Show one decision only if judgment is required.", "Less review load."),
        ("Urgent approval request", "Suppress urgency unless the source establishes a deadline.", "Route to the correct authority threshold.", "Fewer pressure-driven errors."),
        ("Customer data requested for analysis", "Minimize to aggregate metrics and state purpose/retention.", "Block disclosure until permission is evidenced.", "Safer adoption."),
        ("Routine, reversible rename", "Verify authorization, count, and content invariance.", "Pass silently when completion checks succeed.", "Human attention preserved."),
    ], [1.45, 2.1, 1.85, 1.3])
    doc.add_heading("Where the Bouncer stands alone", 2)
    doc.add_paragraph("Bridge-iT Bouncer can be deployed as a skills-only review layer when the immediate pain is duplicated, contradictory, privacy-sensitive, or unsupported agent output. It does not need TAB’s full event-detection layer to add value at a review boundary.")
    doc.add_heading("Where it becomes a TAB add-on", 2)
    doc.add_paragraph("Add the Bouncer Governance Pack to TAB when agent outputs are part of an ongoing workflow and the system must connect hygiene decisions to authority routing, recovery actions, and completion evidence.")

    doc.add_page_break()
    doc.add_heading("4. Product architecture and built features", 1)
    add_table(doc, ["Capability", "What is built", "Value"], [
        ("Pilot-Fit Studio", "Structured input; server-only Lovable Gateway generation; typed diagnosis, intervention, 4-week plan, and KPIs.", "Turns a vague sales conversation into a bounded experiment."),
        ("TAB assistant", "Header trigger; smooth navigation; executive summaries; browser speech synthesis.", "Lets prospects explore at their own pace and accessibility mode."),
        ("Conversational lead capture", "Name, company, email, and challenge collected in chat and prefilled into one shared contact flow.", "Removes duplicate form entry and preserves context."),
        ("Durable contact handoff", "TanStack server function writes private contact records to deployment storage; optional email delivery to danya@dynamicbridge.io.", "No dependency on the visitor sending a mailto draft."),
        ("Investor brief download", "A fixed, one-page PDF served from the site.", "Consistent artifact for partner and investor follow-up."),
        ("Bouncer plugin", "Hygiene Gate and Closure Gate skills; no independent connectors or actions.", "Protects human attention without overstating autonomy."),
    ], [1.35, 3.1, 2.25])
    doc.add_heading("Detect → Coordinate → Recover", 2)
    add_bullets(doc, [
        "Detect: classify the failure using workflow state, source conflicts, missing evidence, permissions, and unresolved ownership.",
        "Coordinate: identify the smallest authorized person, system, or agent that can resolve the exception.",
        "Recover: restore context, complete the missing action, and verify closure against evidence.",
    ])

    doc.add_page_break()
    doc.add_heading("5. Best practices for pilots and production", 1)
    add_table(doc, ["Practice", "Why it matters", "Implementation rule"], [
        ("Start with one seam", "Broad transformation hides causality.", "One workflow, one owner, one baseline."),
        ("Baseline before intervention", "Improvement needs a counterfactual.", "Measure latency, drops, rework, and unresolved exceptions first."),
        ("Minimize data", "More data increases privacy and adoption risk.", "Use aggregate workflow signals; collect personal data only when operationally necessary."),
        ("Encode authority", "Agents should not infer permission from context.", "Store thresholds and approvers as explicit rules."),
        ("Separate reversible from material", "Not every action deserves human attention.", "Silent-pass routine verified work; gate consequential choices."),
        ("Keep source systems authoritative", "A recovery layer must not create shadow truth.", "Link decisions and evidence back to the system of record."),
        ("Verify closure", "Activity is not completion.", "Require downstream evidence appropriate to the action."),
        ("Instrument failure modes", "A pilot must teach the product.", "Log classifications, return reasons, and recovery paths without profiling individuals."),
    ], [1.45, 2.35, 2.9])

    doc.add_page_break()
    doc.add_heading("6. Packaging and pricing hypothesis", 1)
    callout = doc.add_table(rows=1, cols=1)
    set_cell_shading(callout.cell(0, 0), "FFF4E5")
    set_cell_text(callout.cell(0, 0), "These are proposed TAB prices, not reported market facts. Validate willingness to pay in paid pilots before publishing a rate card.", True, "8A4B08", 9)
    doc.add_paragraph()
    add_table(doc, ["Offer", "Benchmark-derived range", "When to sell it", "Includes"], [
        ("4-week design-partner pilot", "$25k–$50k fixed", "One costly, repeated workflow with a committed owner and accessible baseline.", "Workflow map, instrumented seam, 1–3 interventions, weekly readout, evidence pack."),
        ("TAB Core — one environment", "$75k–$180k ARR + usage", "After the pilot proves a measurable recovery loop.", "Detection, routing, continuity, metrics, governance controls, support."),
        ("Bouncer — standalone", "$15k–$40k ARR per controlled workspace", "Review overload is the urgent problem; full workflow instrumentation is premature.", "Hygiene Gate, Closure Gate, policies, eval set, rollout support."),
        ("Bouncer Governance Pack", "$20k–$60k ARR add-on", "TAB already orchestrates mixed human/agent work.", "Agent-output hygiene, decision minimization, authority checks, closure evidence."),
        ("Connector / regulated pack", "$15k–$45k ARR each", "A new system boundary or compliance regime creates material delivery work.", "Connector maintenance, policy mapping, audit controls, testing."),
    ], [1.3, 1.25, 2.1, 2.0])
    doc.add_heading("Pricing logic", 2)
    doc.add_paragraph("Live reference points show multiple accepted enterprise AI buying models: Microsoft 365 Copilot is $30/user/month on annual billing and requires a qualifying plan; Salesforce lists Agentforce at $500 per 100,000 Flex Credits, $2 per conversation, $125/user/month for employee add-ons, and $5/user/month plus usage for company-wide access. TAB should avoid copying seat pricing blindly: its value is the workflow recovered, while usage still matters for cost control.")

    doc.add_page_break()
    doc.add_heading("7. Market context and buyers", 1)
    add_table(doc, ["Published indicator", "Current figure", "What it means for TAB"], [
        ("IDC Worldwide AI and GenAI Spending Guide (2024 V2)", "AI spend: nearly $235B; forecast above $631B by 2028.", "The budget pool is expanding, but TAB must win from workflow reliability, not generic AI spend."),
        ("Gartner agentic AI forecast", "33% of enterprise apps include agentic AI and 15% of day-to-day work decisions are autonomous by 2028.", "More agent actions increase the need for authority, hygiene, and closure."),
        ("Gartner enterprise application exposure", "Up to $234B of enterprise application spend exposed to agentic arbitrage through 2030.", "Per-seat SaaS pricing is under pressure; workflow- or outcome-aligned pricing is credible."),
        ("Gartner 2026 FDE warning", "70% of enterprises predicted to abandon FDE-built agentic AI by 2028 due cost and inability to evolve it.", "TAB must be portable, inspectable, and owned operationally—not a permanent services dependency."),
    ], [2.15, 2.05, 2.5])
    doc.add_heading("Primary buyers", 2)
    add_table(doc, ["Buyer", "Pain", "Entry proof"], [
        ("COO / Head of Operations", "Dropped handoffs, slow exceptions, invisible rework.", "Recovery latency and handoff-drop reduction."),
        ("CIO / Digital Workplace", "AI adoption without cross-tool execution reliability.", "Governed recovery across existing systems."),
        ("Customer Operations leader", "Escalations, onboarding gaps, repeated contacts.", "Faster readiness and fewer unresolved transitions."),
        ("AI transformation / governance lead", "Human review becoming the bottleneck.", "Fewer low-value reviews with authority preserved."),
        ("Risk / Compliance partner", "Unsupported approvals and excess data exposure.", "Explicit permission boundaries and auditable closure."),
    ], [1.7, 2.55, 2.45])
    doc.add_paragraph("Do not present the full AI-spend forecast as TAB’s market size. Use it as context. A defensible serviceable market requires bottom-up counts of target workflows, annual contract value, buyer budget, deployment capacity, and expansion rate.")

    doc.add_page_break()
    doc.add_heading("8. Behind the scenes", 1)
    add_table(doc, ["Layer", "Mechanism", "Failure-handling principle"], [
        ("Signal intake", "Workflow events, source records, agent outputs, and explicit user input.", "Collect the minimum necessary fields; identify the source."),
        ("Breakdown classifier", "Decision vacuum, permission wall, dropped context, ownership ambiguity, review bottleneck, failed handoff.", "Prefer “unknown / needs evidence” over false certainty."),
        ("Policy and authority", "Spend thresholds, data rules, roles, and prior decisions.", "An agent cannot promote itself into an approver."),
        ("Recovery router", "Person, system, or agent best positioned to resolve the exception.", "Route corrections back before consuming human attention."),
        ("Continuity record", "State, prior evidence, unresolved dependencies, and next action.", "Prevent repeated work and context loss."),
        ("Closure proof", "Transaction receipt, state change, acknowledged handoff, or verified outcome.", "Never equate a recommendation or message with execution."),
    ], [1.4, 3.2, 2.1])
    doc.add_heading("Current website stack", 2)
    add_bullets(doc, [
        "TanStack Start and TanStack Router for the full-stack React application and server functions.",
        "React, Tailwind CSS, Radix UI, Lucide, and Framer Motion for the interface.",
        "TanStack AI with the Lovable AI Gateway adapter for server-side pilot generation.",
        "Private Vercel Blob records for durable contact capture; Resend is an optional delivery channel.",
        "Web Speech API for on-device narration; no audio recording is required.",
        "ReportLab and python-docx for controlled investor and internal artifacts.",
    ])
    doc.add_heading("Operational controls", 2)
    add_bullets(doc, [
        "Rate-limit public AI and contact endpoints before high-volume promotion.",
        "Keep model prompts and gateway keys server-side.",
        "Set storage retention and access-review policy for contact requests.",
        "Log model/version, input class, latency, failure code, and output validation result.",
        "Test with missing keys, provider 402/429 responses, malformed model JSON, unsupported speech synthesis, and duplicate form submission.",
    ])

    doc.add_page_break()
    doc.add_heading("9. Evidence register and next decisions", 1)
    doc.add_heading("Sources used for market and pricing context", 2)
    sources = [
        ("S1", "IDC, “Worldwide AI and Generative AI Spending — Industry Outlook” (2024 V2): nearly $235B current market, over $631B by 2028.", "https://www.idc.com/resource-center/blog/idcs-worldwide-ai-and-generative-ai-spending-industry-outlook/"),
        ("S2", "Gartner, “Over 40% of Agentic AI Projects Will Be Canceled by End 2027” (25 Jun 2025): 33% of enterprise apps and 15% of day-to-day decisions by 2028.", "https://www.gartner.com/en/newsroom/press-releases/2025-06-25-gartner-predicts-over-40-percent-of-agentic-ai-projects-will-be-canceled-by-end-of-2027"),
        ("S3", "Gartner, “$234 Billion in Enterprise Application Software Spend Is at Risk from Agentic AI” (1 Jul 2026).", "https://www.gartner.com/en/newsroom/press-releases/2026-07-01-gartner-says-us-dollars-234-billion-in-enterprise-application-software-spend-is-at-risk-from-agentic-artificial-intelligence"),
        ("S4", "Gartner, “70% of Enterprises Will Abandon Agentic AI Built by Vendor FDE by 2028” (29 Sep 2026).", "https://www.gartner.com/en/newsroom/press-releases/2026-09-29-gartner-predicts-70-percent-of-enterprises-will-abandon-agentic-ai-built-by-vendor-forward-deployed-engineering-by-2028"),
        ("S5", "Microsoft, Microsoft 365 Copilot enterprise pricing: $30/user/month, paid yearly; qualifying plan required.", "https://www.microsoft.com/en-us/copilot/pricing/enterprise"),
        ("S6", "Salesforce, Agentforce pricing: $500/100k Flex Credits, $2/conversation, $125/user/month add-on, and $5/user/month plus usage.", "https://www.salesforce.com/agentforce/pricing/"),
        ("S7", "Lovable documentation, AI features for apps and current Gateway model/support lifecycle.", "https://docs.lovable.dev/features/ai"),
        ("S8", "TanStack documentation, Lovable AI Gateway adapter and server functions.", "https://tanstack.com/ai/latest/docs/adapters/lovable"),
    ]
    for key, title, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = True
        p.add_run(f"{key}  ").bold = True
        p.add_run(title + "\n")
        r = p.add_run(url)
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(102, 112, 133)
    doc.add_heading("Decisions to validate in the next 30 days", 2)
    add_bullets(doc, [
        "Which single workflow has enough event volume and cost to justify a paid four-week pilot?",
        "Will buyers accept per-environment pricing, or do procurement systems force a seat/usage hybrid?",
        "Is Bouncer the faster standalone wedge, or does it sell best as governance inside the TAB recovery loop?",
        "Which two connector families recur across qualified prospects?",
        "What evidence threshold converts a pilot into production expansion?",
    ])
    doc.add_paragraph("Document owner: Danya Dotan · Bridge-iT", style="Caption")
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    investor_pdf()
    internal_doc()
    print(PDF_PATH)
    print(DOCX_PATH)
